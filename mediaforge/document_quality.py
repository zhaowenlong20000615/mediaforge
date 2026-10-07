"""Document conversion adapters using established converters, with content checks."""
from pathlib import Path
from io import BytesIO
import base64
import re
import unicodedata
from urllib.parse import urlsplit
from .errors import ForgeError


def convert_word_semantics(source,target,format):
    import mammoth
    import bleach
    from markdownify import markdownify
    from PIL import Image
    from bs4 import BeautifulSoup
    image_count=0
    def inline_image(image):
        nonlocal image_count
        with image.open() as inp:
            raw=inp.read()
        # Turn supported raster images into safe inline PNG; never pass SVG scripts.
        try:
            with Image.open(BytesIO(raw)) as im:
                im.load();out=BytesIO();im.save(out,format='PNG',icc_profile=im.info.get('icc_profile'))
        except Exception:
            raise ForgeError('unsupported_embedded_image','文档中包含不能可靠转换的图片。','改为导出 PDF 保留外观，或先将嵌入图片替换为 PNG/JPEG。')
        image_count+=1
        return {'src':'data:image/png;base64,'+base64.b64encode(out.getvalue()).decode()}
    try:
        with Path(source).open('rb') as inp:
            result=mammoth.convert_to_html(inp,convert_image=mammoth.images.img_element(inline_image),external_file_access=False,include_embedded_style_map=False)
    except ForgeError:raise
    except Exception:raise ForgeError('document_conversion','Word 内容无法可靠转换。','请重新保存为 DOCX，或尝试 PDF 导出。')
    def attribute(tag,name,value):
        if tag=='img':return name=='alt' or name=='src' and value.startswith('data:image/png;base64,')
        if tag=='a' and name=='href':return urlsplit(value).scheme.lower() in {'','http','https','mailto'}
        return name in {'id','colspan','rowspan'}
    clean=bleach.clean(result.value,tags={'p','h1','h2','h3','h4','h5','h6','ul','ol','li','strong','em','u','s','sup','sub','br','hr','a','img','table','thead','tbody','tr','td','th','blockquote','pre','code'},attributes=attribute,protocols={'http','https','mailto','data'},strip=True)
    soup=BeautifulSoup(clean,'html.parser')
    if not soup.get_text(strip=True) and not image_count:raise ForgeError('no_content','文档没有可转换的内容。')
    if format=='html':
        style='body{max-width:900px;margin:3rem auto;padding:0 1.5rem;font-family:Arial,"Noto Sans CJK SC",sans-serif;line-height:1.6;color:#20272b}img{max-width:100%;height:auto}table{border-collapse:collapse;width:100%}td,th{border:1px solid #bbb;padding:.5rem}h1,h2,h3{line-height:1.25}'
        text='<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; img-src data:; style-src \'unsafe-inline\'; base-uri \'none\'"><title>Converted document</title><style>'+style+'</style></head><body>'+clean+'</body></html>'
    elif format=='md':
        text=markdownify(clean,heading_style='ATX',bullets='-',table_infer_header=True)
    else:text=soup.get_text('\n',strip=True)
    Path(target).write_text(text,encoding='utf-8')
    notes=[]
    if result.messages:notes.append('部分 Word 样式未映射，请与原文核对。')
    from docx import Document
    document=Document(source)
    if any(p.text.strip() for section in document.sections for part in [section.header,section.footer,section.first_page_header,section.first_page_footer,section.even_page_header,section.even_page_footer] for p in part.paragraphs):
        notes.append('此格式导出正文；原文含页眉或页脚，完整外观请导出 PDF。')
    return {'engine':'mammoth','embedded_images':image_count,'tables':len(soup.select('table')),'headings':len(soup.select('h1,h2,h3,h4,h5,h6')),'mode':'semantic_content','conversion_notices':len(result.messages),'quality_notes':notes}


def pdf_to_docx_layout(source,target):
    """Pinned pdf2docx adapter: unmodified upstream, isolated by the job process."""
    from pdf2docx import Converter
    from docx import Document
    import pymupdf
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    converter=Converter(str(source))
    # Borderless-table inference can turn ordinary bullet lists into narrow,
    # fixed-height cells. Explicit bordered tables remain editable tables.
    try:converter.convert(str(target),multi_processing=False,ignore_page_error=False,parse_stream_table=False)
    finally:converter.close()
    doc=Document(target)
    def repair(container):
        for para in container.paragraphs:
            for run in para.runs:
                if '\uf0b7' in run.text:
                    run.text=run.text.replace('\uf0b7','•');run.font.name='Arial'
        for table in container.tables:
            for row in table.rows:
                if row.height_rule==WD_ROW_HEIGHT_RULE.EXACTLY:row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
                for cell in row.cells:repair(cell)
    repair(doc);doc.save(target)
    def text_of(container):
        values=[p.text for p in container.paragraphs]
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:values.extend(text_of(cell))
        return values
    text='\n'.join(text_of(doc))
    with pymupdf.open(source) as pdf:
        source_text='\n'.join(page.get_text() for page in pdf)
        images=sum(len(page.get_images()) for page in pdf)
        pages=len(pdf)
    # Compare normalized character inventory; reading-order and visual layout are
    # verified by rendering in acceptance fixtures, not falsely inferred here.
    from collections import Counter
    normalize=lambda s:Counter(normalize_document_text(s))
    expected=normalize(source_text);actual=normalize(text)
    coverage=sum((expected&actual).values())/max(1,sum(expected.values()))
    if source_text.strip() and coverage<.95:
        raise ForgeError('document_content_loss','版式转换丢失了过多文字，结果未交付。','请改为 text 方式提取文字，或先对扫描 PDF 进行 OCR。')
    if images and not doc.inline_shapes:
        raise ForgeError('document_images_missing','检测到原 PDF 的图片没有保留。','请使用原 PDF 或其他保真格式；避免误交付缺图的 Word。')
    return {'engine':'pdf2docx','source_pages':pages,'mode':'layout','text_character_coverage':round(coverage,5),'embedded_images':len(doc.inline_shapes),'tables':len(doc.tables),'visual_review_required':True}


def normalize_document_text(text):
    return re.sub(r'\s+','',unicodedata.normalize('NFKC',text.replace('\uf0b7','•')))


def verify_document_render(source,rendered):
    """Check visible PDF text after Office layout, catching clipped DOCX text."""
    from collections import Counter
    import pymupdf
    def content(path):
        with pymupdf.open(path) as doc:return '\n'.join(page.get_text() for page in doc),len(doc)
    before,_=content(source);after,pages=content(rendered)
    expected=Counter(normalize_document_text(before));actual=Counter(normalize_document_text(after))
    coverage=sum((expected&actual).values())/max(1,sum(expected.values()))
    numbers=lambda text:Counter(re.findall(r'\d+(?:[.,-]\d+)*',unicodedata.normalize('NFKC',text)))
    if coverage<.98 or numbers(before)-numbers(after):
        raise ForgeError('document_layout_loss','Word 排版后出现文字或数字缺失，结果未交付。','改用 text 方式获取文字，或保留原 PDF；复杂版式需要专门排版。')
    return {'rendered_text_coverage':round(coverage,5),'rendered_numbers_preserved':True,'rendered_pages':pages}
