from contextlib import asynccontextmanager
from pathlib import Path
from collections import defaultdict,deque
from urllib.parse import urlsplit
import asyncio
import io
import json
import os
import secrets
import time
from fastapi import FastAPI,Request,Depends,Query
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse,FileResponse,Response
from pydantic import BaseModel,Field,ConfigDict
from . import __version__
from .config import Settings
from .store import Store,clean_name
from .engine import Engine
from .errors import ForgeError
from .catalog import OPS
from .dependencies import checks,catalog_status,ocr_languages


class TaskInput(BaseModel):
    model_config=ConfigDict(extra='forbid')
    tool:str
    file_ids:list[str]
    params:dict=Field(default_factory=dict)
    idempotency_key:str|None=None
    group:str=''
    timeout:int|None=None
    retries:int=Field(default=0,ge=0,le=3)


async def small_json(request):
    body=bytearray()
    async for chunk in request.stream():
        body.extend(chunk)
        if len(body)>65536:raise ForgeError('request_limit','请求体超过大小限制。','请减少参数或文件数量。',413)
    try:
        value=json.loads(body)
        if not isinstance(value,dict):raise ValueError()
        return value
    except (ValueError,UnicodeDecodeError):raise ForgeError('invalid_json','请求格式不正确。','请使用 JSON 对象。')


def create_app(settings=None,start_worker=True):
    settings=settings or Settings.load(); store=Store(settings); engine=Engine(store)
    attempts=defaultdict(deque)
    web=Path(__file__).parent/'web'

    @asynccontextmanager
    async def lifespan(app):
        if start_worker:engine.start()
        yield
        if start_worker:await asyncio.to_thread(engine.stop)

    app=FastAPI(title='MediaForge',version=__version__,lifespan=lifespan,docs_url=None,redoc_url=None,openapi_url=None)
    app.state.store=store;app.state.engine=engine

    def origin_ok(request):
        origin=request.headers.get('origin')
        if not origin:return True
        configured=os.getenv('MEDIAFORGE_PUBLIC_URL','')
        allowed={str(request.base_url).rstrip('/')}
        if configured:
            parsed=urlsplit(configured);allowed.add(parsed.scheme+'://'+parsed.netloc)
        return origin.rstrip('/') in allowed

    @app.middleware('http')
    async def boundary(request,call_next):
        request_id=secrets.token_hex(8);request.state.request_id=request_id
        if request.method not in {'GET','HEAD','OPTIONS'} and not origin_ok(request):
            return JSONResponse({'error':{'code':'invalid_origin','message':'请求来源不受信任。','action':'从工作台所在地址重新打开页面。'}},status_code=403)
        if request.method not in {'GET','HEAD'} and request.url.path!='/api/files':
            try:
                if int(request.headers.get('content-length','0'))>65536:raise ValueError()
            except ValueError:
                return JSONResponse({'error':{'code':'request_limit','message':'请求体过大。'}},status_code=413)
        response=await call_next(request)
        response.headers['X-Content-Type-Options']='nosniff'
        response.headers['Referrer-Policy']='no-referrer'
        response.headers['X-Frame-Options']='SAMEORIGIN'
        response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data: blob:; media-src 'self' blob:; frame-src 'self' blob:; object-src 'none'; base-uri 'self'; frame-ancestors 'self'; form-action 'self'"
        response.headers['X-Request-ID']=request_id
        if request.url.path.startswith('/api/'):response.headers['Cache-Control']='no-store'
        # Method and route template only: never parameters, headers, filenames or body.
        route=request.scope.get('route');route_path=getattr(route,'path','unmatched')
        print(json.dumps({'time':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),'project_id':'local.mediaforge','request_id':request_id,'task_id':None,'method':request.method,'route':route_path,'status':response.status_code}),flush=True)
        return response

    @app.exception_handler(ForgeError)
    async def forge_error(request,e):
        return JSONResponse({'error':e.payload(),'request_id':getattr(request.state,'request_id','')},status_code=e.status)

    @app.exception_handler(RequestValidationError)
    async def validation_error(request,e):
        return JSONResponse({'error':{'code':'invalid_request','message':'请求参数不符合接口要求。','action':'检查工具参数和文件 ID；不要提交服务器路径。'}},status_code=422)

    @app.exception_handler(Exception)
    async def internal_error(request,e):
        return JSONResponse({'error':{'code':'internal_error','message':'服务处理异常。','action':'保留任务编号并重试。'},'request_id':getattr(request.state,'request_id','')},status_code=500)

    def owner(request:Request):
        auth=request.headers.get('authorization','')
        token=auth[7:] if auth.startswith('Bearer ') else request.cookies.get('mf_session','')
        value=store.authenticate(token)
        if not value:raise ForgeError('unauthorized','请先连接工作区。','输入访问令牌后继续。',401)
        return value

    @app.get('/healthz')
    def health():return {'status':'ok','project_id':'local.mediaforge'}

    @app.get('/readyz')
    def ready():
        ready=engine.ready if start_worker else True
        status=checks(); missing=[x['id'] for x in status if not x['available']]
        return JSONResponse({'status':('ready_degraded' if missing else 'ready') if ready else 'not_ready','worker_ready':ready,'optional_or_external_dependencies_missing':missing,'meaning':'服务就绪不代表所有工具或模型已配置；工具目录给出每项可用性。'},status_code=200 if ready else 503)

    @app.get('/version')
    def version():
        release=Path(os.getenv('MEDIAFORGE_RELEASE_FILE',str(Path(__file__).parent.parent/'release.json')))
        try:meta=json.loads(release.read_text())
        except (ValueError,OSError):meta={}
        return {'name':'mediaforge','version':__version__,'git_commit':meta.get('git_commit','development'),'artifact_sha256':meta.get('artifact_sha256'), 'built_at':meta.get('built_at')}

    @app.post('/api/session')
    async def login(request:Request):
        ip=request.client.host if request.client else 'local';now=time.monotonic();q=attempts[ip]
        while q and now-q[0]>60:q.popleft()
        if len(q)>=10:raise ForgeError('login_rate_limit','尝试过于频繁。','一分钟后重试。',429)
        q.append(now)
        data=await small_json(request)
        if not isinstance(data,dict) or not isinstance(data.get('token'),str):raise ForgeError('invalid_token','请输入访问令牌。')
        token,value=store.session(data['token'])
        response=JSONResponse({'owner':value,'expires_in':43200})
        response.set_cookie('mf_session',token,max_age=43200,httponly=True,secure=settings.secure_cookie,samesite='strict',path=os.getenv('MEDIAFORGE_COOKIE_PATH','/'))
        return response

    @app.delete('/api/session')
    def logout(request:Request,who=Depends(owner)):
        store.revoke(request.cookies.get('mf_session',''))
        res=JSONResponse({'signed_out':True});res.delete_cookie('mf_session',path=os.getenv('MEDIAFORGE_COOKIE_PATH','/'));return res

    @app.get('/api/me')
    def me(who=Depends(owner)):return {'owner':who,'storage':store.storage(who)}

    @app.get('/api/tools')
    def tools(who=Depends(owner)):return {'tools':catalog_status(OPS)}

    @app.get('/api/doctor')
    def doctor(who=Depends(owner)):return {'checks':checks(),'ocr_languages':ocr_languages(),'storage':store.storage(who),'server_platform':sys_platform(),'worker_ready':engine.ready}

    @app.post('/api/files',status_code=201)
    async def upload(request:Request,name:str=Query(min_length=1,max_length=220),who=Depends(owner)):
        name=clean_name(name)
        try:size=int(request.headers.get('content-length','0'))
        except ValueError:raise ForgeError('invalid_length','文件大小无效。')
        reservation=store.reserve(who,size);temp=store.root/'staging'/reservation
        try:
            count=0
            with temp.open('wb') as f:
                async for chunk in request.stream():
                    count+=len(chunk)
                    if count>size or count>settings.max_file_bytes:raise ForgeError('file_too_large','上传超过大小限制。',status=413)
                    f.write(chunk)
            if count!=size:raise ForgeError('upload_incomplete','文件上传未完成。','保留原文件并重新上传。')
            temp.chmod(0o600)
            return await asyncio.to_thread(store.register_upload,who,name,temp,reservation)
        finally:
            temp.unlink(missing_ok=True);store.release_reservation(reservation)

    @app.get('/api/files')
    def file_list(who=Depends(owner),limit:int=Query(30,ge=1,le=100),offset:int=Query(0,ge=0)):
        return {'files':store.file_list(who,limit,offset)}

    @app.get('/api/files/{id}')
    def file_meta(id:str,who=Depends(owner)):return store.file(who,id)

    @app.get('/api/files/{id}/content')
    def content(id:str,inline:bool=False,who=Depends(owner)):
        f=store.file(who,id,True)
        safe_inline=f['kind'] in {'image','video','audio','pdf'}
        return FileResponse(f['path'],media_type=f['mime'] if safe_inline else 'application/octet-stream',filename=f['name'],
                            content_disposition_type='inline' if inline and safe_inline else 'attachment')

    @app.get('/api/files/{id}/preview')
    def preview(id:str,who=Depends(owner)):
        f=store.file(who,id,True);result={'kind':f['kind'],'file_id':id,'name':f['name'],'quality_review':'required'}
        if f['kind'] in {'text','subtitle'}:
            with open(f['path'],'rb') as inp:data=inp.read(64001)
            result.update(text=data[:64000].decode('utf-8-sig','replace'),truncated=len(data)>64000)
        elif f['kind']=='docx':
            from docx import Document
            doc=Document(f['path']);text='\n'.join(p.text for p in doc.paragraphs)
            for t in doc.tables:
                text+='\n'+'\n'.join('\t'.join(c.text for c in r.cells) for r in t.rows)
            result.update(text=text[:64000],truncated=len(text)>64000,layout_preview=False)
        elif f['kind'] in {'image','video','audio','pdf'}:result['inline']=True
        else:result['download_only']=True
        return result

    @app.get('/api/files/{id}/thumbnail')
    def thumbnail(id:str,who=Depends(owner)):
        from PIL import Image,ImageOps
        f=store.file(who,id,True)
        if f['kind']!='image':raise ForgeError('no_thumbnail','此文件没有图片缩略图。',status=404)
        with Image.open(f['path']) as src:
            im=ImageOps.exif_transpose(src);im.thumbnail((320,240));out=io.BytesIO();im.convert('RGB').save(out,format='JPEG',quality=75)
        return Response(out.getvalue(),media_type='image/jpeg')

    @app.post('/api/tasks',status_code=202)
    async def submit(request:Request,who=Depends(owner)):
        try:data=TaskInput.model_validate(await small_json(request))
        except Exception as e:
            if isinstance(e,ForgeError):raise
            raise ForgeError('invalid_request','任务参数无效。','使用文件 ID，并检查工具参数。',422)
        return store.submit(who,data.tool,data.file_ids,data.params,data.idempotency_key,data.group,data.timeout,data.retries)

    @app.get('/api/tasks')
    def tasks(who=Depends(owner),status:str='',search:str=Query('',max_length=200),group:str='',limit:int=Query(30,ge=1,le=100),offset:int=Query(0,ge=0)):
        return store.tasks(who,status,search,limit,offset,group)

    @app.get('/api/tasks/{id}')
    def task(id:str,who=Depends(owner)):return store.task(who,id)

    @app.get('/api/tasks/{id}/logs')
    def logs(id:str,who=Depends(owner)):return {'events':store.logs(who,id)}

    @app.post('/api/tasks/{id}/cancel')
    def cancel(id:str,who=Depends(owner)):return store.cancel(who,id)

    @app.post('/api/tasks/{id}/resume')
    def resume(id:str,who=Depends(owner)):return store.resume(who,id)

    @app.post('/api/exports',status_code=202)
    async def export(request:Request,who=Depends(owner)):
        data=await small_json(request);ids=data.get('task_ids',[])
        if not isinstance(ids,list) or not 1<=len(ids)<=50:raise ForgeError('invalid_export','选择 1–50 个任务。')
        files=[]
        for tid in ids:files.extend(store.task(who,tid,False)['output_ids'])
        if not files:raise ForgeError('no_results','所选任务没有可导出的结果。')
        return store.submit(who,'export-results',list(dict.fromkeys(files)),{},data.get('idempotency_key'),group='结果导出')

    @app.get('/api/templates')
    def templates(who=Depends(owner)):return {'templates':store.templates(who)}

    @app.post('/api/templates',status_code=201)
    async def template(request:Request,who=Depends(owner)):
        d=await small_json(request)
        if not isinstance(d.get('name'),str) or not d['name'].strip():raise ForgeError('template_name','请输入模板名称。')
        return store.save_template(who,d['name'],d.get('tool'),d.get('params',{}))

    @app.delete('/api/templates/{id}')
    def delete_template(id:str,who=Depends(owner)):
        with store.tx() as con:con.execute('DELETE FROM templates WHERE id=? AND owner=?',(id,who))
        return {'deleted':True}

    @app.post('/api/cleanup')
    async def cleanup(request:Request,who=Depends(owner)):
        d=await small_json(request);days=d.get('days',settings.retention_days)
        if not isinstance(days,int) or days<0:raise ForgeError('invalid_days','保留天数必须为非负整数。')
        if type(d.get('dry_run',True)) is not bool:raise ForgeError('invalid_request','dry_run 必须是布尔值。')
        return await asyncio.to_thread(store.cleanup,who,days,d.get('dry_run',True))

    @app.post('/api/workspaces',status_code=201)
    async def workspace(request:Request,who=Depends(owner)):
        d=await small_json(request)
        return store.create_workspace(who,str(d.get('name','新工作区')))

    @app.get('/')
    def index():return FileResponse(web/'index.html',media_type='text/html')

    @app.get('/assets/{name}')
    def asset(name:str):
        allowed={'app.js':'application/javascript','style.css':'text/css'}
        if name not in allowed:raise ForgeError('not_found','资源不存在。',status=404)
        return FileResponse(web/name,media_type=allowed[name])

    return app


def sys_platform():
    import platform
    return {'system':platform.system(),'architecture':platform.machine()}


def main():
    import uvicorn
    settings=Settings.load()
    uvicorn.run(create_app(settings),host=settings.host,port=settings.port,uds=os.getenv('MEDIAFORGE_UNIX_SOCKET'),access_log=False,proxy_headers=True,forwarded_allow_ips='127.0.0.1')

if __name__=='__main__':main()
