from pathlib import Path
import shutil
import subprocess
import pytest
from PIL import Image,ImageDraw
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from pypdf import PdfReader,PdfWriter
from mediaforge.config import Settings
from mediaforge.store import Store
from mediaforge.engine import Engine


@pytest.fixture(scope='session')
def samples(tmp_path_factory):
    root=tmp_path_factory.mktemp('samples')
    im=Image.new('RGB',(800,600),'white');d=ImageDraw.Draw(im);d.rectangle((50,100,750,500),fill=(50,90,130));d.ellipse((200,150,600,500),fill=(235,150,70));im.save(root/'图片 sample.png')
    mask=Image.new('L',(800,600),0);ImageDraw.Draw(mask).rectangle((300,280,330,310),fill=255);mask.save(root/'mask.png')
    c=canvas.Canvas(str(root/'document.pdf'),pagesize=A4)
    for page in range(3):
        c.drawString(60,780,'MEDIAFORGE TEST PAGE '+str(page+1))
        if page==0:
            for x in [60,200,340]:c.line(x,600,x,700)
            for y in [600,650,700]:c.line(60,y,340,y)
            c.drawString(70,670,'Item');c.drawString(210,670,'Amount');c.drawString(70,620,'Apple');c.drawString(210,620,'42')
        c.showPage()
    c.save()
    writer=PdfWriter()
    for page in PdfReader(root/'document.pdf').pages:writer.add_page(page)
    writer.encrypt('audit-password',algorithm='AES-256')
    with (root/'protected.pdf').open('wb') as out:writer.write(out)
    (root/'broken.pdf').write_bytes(b'%PDF-1.7\nnot a valid document')
    doc=Document();p=doc.add_paragraph();p.add_run('Hello ').bold=True;p.add_run('World');doc.add_paragraph('MediaForge document text')
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Name';table.cell(0,1).text='Count';table.cell(1,0).text='World';table.cell(1,1).text='7';doc.sections[0].header.paragraphs[0].text='World header';doc.save(root/'input.docx')
    (root/'subtitle.srt').write_text('1\n00:00:00,000 --> 00:00:01,000\nHello MediaForge\n\n2\n00:00:01,000 --> 00:00:02,000\nUnicode 世界\n',encoding='utf-8')
    (root/'plain.txt').write_text('not an image',encoding='utf-8')
    if shutil.which('ffmpeg'):
        def ff(args):subprocess.run(['ffmpeg','-y','-v','error',*map(str,args)],check=True,capture_output=True,timeout=30)
        ff(['-f','lavfi','-i','sine=frequency=440:duration=2',root/'audio.wav'])
        ff(['-f','lavfi','-i','color=c=blue:s=320x240:d=2','-i',root/'audio.wav','-c:v','libx264','-c:a','aac','-shortest',root/'video.mp4'])
        ff(['-f','lavfi','-i','color=c=red:s=320x240:d=1','-c:v','libx264',root/'silent.mp4'])
        ff(['-i',root/'video.mp4','-i',root/'audio.wav','-i',root/'subtitle.srt','-map','0:v','-map','0:a','-map','1:a','-map','2:s','-c:v','copy','-c:a','aac','-c:s','mov_text',root/'tracks.mp4'])
    return root


@pytest.fixture
def store(tmp_path):return Store(Settings(data=tmp_path/'data',concurrency=2))

@pytest.fixture
def running(store):
    engine=Engine(store);engine.start()
    yield store,engine
    engine.stop()
