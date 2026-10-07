"""Output fidelity tests use actual codecs/converters, not status-only assertions."""
from pathlib import Path
from io import BytesIO
import json
import subprocess
import pytest
import numpy as np
from PIL import Image,ImageCms,ImageDraw
from pypdf import PdfReader
from docx import Document
from docx.shared import Inches
from mediaforge.dependencies import verified_binary
from mediaforge.files import probe
from mediaforge.subtitle_quality import cues_from_words,display_width
from tests.test_core import job,output_path


def test_image_pdf_embeds_original_pixels_and_jpeg_bytes(running,tmp_path):
    store,_=running
    a=np.random.default_rng(1234).integers(0,256,(120,180,3),dtype=np.uint8)
    png=tmp_path/'detail.png';Image.fromarray(a).save(png)
    jpeg=tmp_path/'photo.jpg';Image.fromarray(a).save(jpeg,quality=96,subsampling=0)
    t=job(store,'image-pdf',[png,jpeg]);assert t['status']=='succeeded',t['error']
    pdf=PdfReader(output_path(store,t));assert len(pdf.pages)==2
    assert np.array_equal(np.array(pdf.pages[0].images[0].image.convert('RGB')),a)
    # pypdf's image convenience accessor decodes/re-encodes JPEG. Inspect the
    # actual DCT stream to verify the PDF embeds the original compressed bytes.
    xobject=next(iter(pdf.pages[1]['/Resources']['/XObject'].values())).get_object()
    assert xobject['/Filter']=='/DCTDecode' and xobject.get_data()==jpeg.read_bytes()
    assert t['outputs'][0]['verification']['image_recompression'] is False


def test_image_default_webp_is_lossless_and_icc_retained(running,tmp_path):
    store,_=running
    a=np.random.default_rng(8).integers(0,256,(120,180,3),dtype=np.uint8)
    im=Image.fromarray(a);p=tmp_path/'color.png';icc=ImageCms.ImageCmsProfile(ImageCms.createProfile('sRGB')).tobytes();im.save(p,icc_profile=icc)
    t=job(store,'image-process',[p],{'format':'webp'});assert t['status']=='succeeded',t['error']
    out=Image.open(output_path(store,t));assert np.array_equal(np.array(out.convert('RGB')),a)
    assert out.info['icc_profile']==icc


def test_word_semantic_export_keeps_images_lists_tables_styles(running,tmp_path):
    store,_=running
    image=tmp_path/'figure.png';Image.new('RGB',(80,60),'red').save(image)
    d=Document();d.add_heading('Quality Title',0);d.add_heading('Important section',1);d.add_paragraph('First item',style='List Bullet');p=d.add_paragraph();p.add_run('Important').bold=True
    table=d.add_table(rows=2,cols=2);table.cell(0,0).text='Amount';table.cell(1,0).text='1234.56';d.add_picture(str(image),width=Inches(1));source=tmp_path/'document.docx';d.save(source)
    for fmt in ['html','md']:
        t=job(store,'office-convert',[source],{'format':fmt});assert t['status']=='succeeded',t['error']
        text=output_path(store,t).read_text();assert '1234.56' in text and 'data:image/png;base64,' in text
        if fmt=='html':assert '<h1>Quality Title</h1>' in text and '<h2>Important section</h2>' in text and '<strong>' in text and '<ul>' in text and '<table>' in text
        else:assert '# Quality Title' in text and '**Important**' in text and '- First item' in text
        assert t['outputs'][0]['verification']['embedded_images']==1


def test_pdf_word_layout_preserves_table_image_and_values(running,tmp_path):
    from reportlab.pdfgen import canvas
    image=tmp_path/'image.png';Image.new('RGB',(80,60),'blue').save(image)
    source=tmp_path/'layout.pdf';c=canvas.Canvas(str(source));c.setFont('Helvetica-Bold',24);c.drawString(50,740,'QUALITY DOCUMENT')
    for x in [50,230,410]:c.line(x,520,x,640)
    for y in [520,580,640]:c.line(50,y,410,y)
    c.setFont('Helvetica',12);c.drawString(60,610,'Invoice');c.drawString(240,610,'1234.56');c.drawString(60,550,'Total');c.drawString(240,550,'1234.56')
    c.drawImage(str(image),70,250,width=240,height=180);c.save()
    store,_=running;t=job(store,'pdf-word',[source]);assert t['status']=='succeeded',t['error']
    d=Document(output_path(store,t));assert len(d.tables)>=1 and len(d.inline_shapes)>=1
    assert any('1234.56' in cell.text for table in d.tables for row in table.rows for cell in row.cells)
    assert t['outputs'][0]['verification']['text_character_coverage']>=.95
    assert len(t['outputs'])==2 and t['outputs'][1]['verification']['role']=='layout_preview'
    assert t['outputs'][0]['verification']['rendered_numbers_preserved'] is True


def test_word_pdf_word_roundtrip_keeps_visible_list_and_amounts(running,tmp_path):
    store,_=running
    d=Document();d.add_heading('Quality Brief',0)
    for value in ['First delivery item','Second delivery item']:d.add_paragraph(value,style='List Bullet')
    table=d.add_table(rows=2,cols=2);table.style='Table Grid';table.cell(0,0).text='Amount';table.cell(0,1).text='120.50';table.cell(1,0).text='Total';table.cell(1,1).text='120.50'
    p=tmp_path/'brief.docx';d.save(p)
    t=job(store,'office-convert',[p],{'format':'pdf'});assert t['status']=='succeeded',t['error']
    t=job(store,'pdf-word',[output_path(store,t)]);assert t['status']=='succeeded',t['error']
    visible=' '.join(page.extract_text() for page in PdfReader(output_path(store,t,1)).pages)
    assert all(value in visible for value in ['First delivery item','Second delivery item','120.50'])


def test_rendered_document_rejects_changed_amount_even_with_complete_body(tmp_path):
    from reportlab.pdfgen import canvas
    from mediaforge.document_quality import verify_document_render
    from mediaforge.errors import ForgeError
    for name,amount in [('source','100.01'),('rendered','100.02')]:
        c=canvas.Canvas(str(tmp_path/(name+'.pdf')))
        for i in range(30):c.drawString(30,780-i*20,'Reference document text should survive unchanged.')
        c.drawString(30,100,'Amount: '+amount);c.save()
    with pytest.raises(ForgeError,match='缺失'):verify_document_render(tmp_path/'source.pdf',tmp_path/'rendered.pdf')


def test_subtitle_cues_are_readable_and_keep_all_words():
    text='欢迎使用媒体工具箱请检查每一份文档和每一张图片的输出质量保持画面清晰并且字幕准确'
    words=[{'text':char,'start':i*.3,'end':i*.3+.25} for i,char in enumerate(text)]
    cues=cues_from_words(words,28)
    assert ''.join(c['text'].replace('\n','') for c in cues)==text
    assert all(len(c['text'].splitlines())<=2 and all(display_width(line)<=28 for line in c['text'].splitlines()) for c in cues)
    assert all(c['end']-c['start']<=6 for c in cues)
    assert all(a['end']<=b['start'] for a,b in zip(cues,cues[1:]))


def test_subtitle_english_words_long_identifiers_and_pauses():
    values=[' Acknowledgement',' characteristics',' transformation',' identifier012345678901234567890123456789',' end.']
    words=[{'text':word,'start':i*8,'end':i*8+7} for i,word in enumerate(values)]
    cues=cues_from_words(words,28)
    assert ''.join(''.join(c['text'].split()) for c in cues)==''.join(''.join(w.split()) for w in values)
    assert all(len(c['text'].splitlines())<=2 and all(display_width(line)<=28 for line in c['text'].splitlines()) for c in cues)
    assert all(c['end']-c['start']<=6 for c in cues)
    assert all(a['end']<=b['start'] for a,b in zip(cues,cues[1:]))


def test_palette_transparency_and_lab_color_conversion(running,tmp_path):
    store,_=running
    p=tmp_path/'palette.png';im=Image.new('P',(40,40));im.putpalette([255,0,0,0,0,255]+[0]*762);im.putpixel((20,20),1);im.save(p,transparency=0)
    t=job(store,'image-process',[p],{'format':'jpeg'});assert t['status']=='succeeded',t['error']
    assert all(x>245 for x in Image.open(output_path(store,t)).getpixel((0,0)))
    srgb=ImageCms.createProfile('sRGB');lab=ImageCms.ImageCmsProfile(ImageCms.createProfile('LAB'));original=Image.new('RGB',(40,40),(180,70,40))
    source=ImageCms.profileToProfile(original,srgb,lab,outputMode='LAB');p=tmp_path/'lab.tiff';source.save(p,icc_profile=lab.tobytes())
    t=job(store,'image-process',[p],{'format':'png'});assert t['status']=='succeeded',t['error']
    out=Image.open(output_path(store,t));assert out.mode=='RGB'
    assert max(abs(a-b) for a,b in zip(out.getpixel((0,0)),original.getpixel((0,0))))<4
    assert t['outputs'][0]['verification']['color_converted_to_rgb'] is True


def test_merge_preserves_geometry_rate_and_decoded_frames(running,tmp_path):
    ff=verified_binary('ffmpeg')
    if not ff:pytest.skip('FFmpeg unavailable')
    source=tmp_path/'clip.mp4'
    subprocess.run([ff,'-v','error','-y','-f','lavfi','-i','testsrc2=size=640x360:rate=60:duration=1','-f','lavfi','-i','sine=frequency=440:duration=1','-c:v','libx264','-crf','18','-pix_fmt','yuv420p','-c:a','aac','-shortest',str(source)],check=True,capture_output=True)
    store,_=running;t=job(store,'media-merge',[source,source],{'kind':'video'});assert t['status']=='succeeded',t['error']
    meta=probe(output_path(store,t));v=next(x for x in meta['streams'] if x['codec_type']=='video')
    assert (v['width'],v['height'])==(640,360) and v['avg_frame_rate']=='60/1'
    assert t['outputs'][0]['verification']['stream_copy'] is True
    def frame(path):return subprocess.check_output([ff,'-v','error','-i',str(path),'-frames:v','1','-f','rawvideo','-pix_fmt','rgb24','-'])
    assert frame(source)==frame(output_path(store,t))


def test_merge_preserves_phone_portrait_orientation(running,tmp_path):
    ff=verified_binary('ffmpeg')
    if not ff:pytest.skip('FFmpeg unavailable')
    source=tmp_path/'landscape.mp4';rotated=tmp_path/'phone.mp4'
    subprocess.run([ff,'-v','error','-y','-f','lavfi','-i','testsrc2=size=320x240:rate=25:duration=1','-c:v','libx264','-pix_fmt','yuv420p',str(source)],check=True,capture_output=True)
    help_text=subprocess.run([ff,'-hide_banner','-h','full'],check=True,capture_output=True,text=True).stdout
    args=['-display_rotation','90','-i',str(source),'-c','copy'] if '-display_rotation' in help_text else ['-i',str(source),'-c','copy','-metadata:s:v:0','rotate=90']
    subprocess.run([ff,'-v','error','-y',*args,str(rotated)],check=True,capture_output=True)
    assert probe(rotated)['streams'][0]['rotation']==90
    store,_=running;t=job(store,'media-merge',[rotated,source],{'kind':'video'});assert t['status']=='succeeded',t['error']
    v=next(x for x in probe(output_path(store,t))['streams'] if x['codec_type']=='video')
    assert (v['width'],v['height'],v['rotation'])==(240,320,0)
    def frame(path):return np.frombuffer(subprocess.check_output([ff,'-v','error','-i',str(path),'-frames:v','1','-f','rawvideo','-pix_fmt','rgb24','-']),dtype=np.uint8).astype(float)
    assert np.mean((frame(rotated)-frame(output_path(store,t)))**2)<50


def test_jpeg_defaults_preserve_chroma_detail(running,tmp_path):
    from skimage.metrics import structural_similarity
    image=Image.new('RGB',(400,200),'white');draw=ImageDraw.Draw(image)
    for x in range(0,400,4):draw.line((x,0,x,199),fill=(20,20,255),width=1)
    source=tmp_path/'detail.png';image.save(source)
    old=BytesIO();image.save(old,format='JPEG',quality=85)
    store,_=running;t=job(store,'image-process',[source],{'format':'jpeg'});assert t['status']=='succeeded',t['error']
    ref=np.array(image);before=structural_similarity(ref,np.array(Image.open(BytesIO(old.getvalue()))),channel_axis=2,data_range=255)
    after=structural_similarity(ref,np.array(Image.open(output_path(store,t))),channel_axis=2,data_range=255)
    assert after>.97 and after>before+.05,(before,after)


def test_detailed_jpeg_does_not_overflow_encoder_buffer(running,tmp_path):
    store,_=running
    a=np.random.default_rng(91).integers(0,256,(512,1024,3),dtype=np.uint8)
    source=tmp_path/'complex.png';Image.fromarray(a).save(source)
    t=job(store,'image-process',[source],{'format':'jpeg'})
    assert t['status']=='succeeded',t['error']
    assert Image.open(output_path(store,t)).size==(1024,512)
